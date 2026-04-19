import openpyxl
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "東京近郊お茶農園リスト"

# --- Colors ---
header_fill   = PatternFill("solid", fgColor="2E7D32")  # dark green
subhdr_fill   = PatternFill("solid", fgColor="A5D6A7")  # light green
row_odd_fill  = PatternFill("solid", fgColor="F1F8E9")
row_even_fill = PatternFill("solid", fgColor="FFFFFF")
thin = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin")
)

# --- Column definitions ---
columns = [
    ("農園名",        18),
    ("都道府県",       9),
    ("住所",         35),
    ("電話番号",      14),
    ("メール",        28),
    ("公式URL",      38),
    ("都心からの距離",  14),
    ("ほうじ茶・焙煎",  22),
    ("体験プログラム",  38),
    ("ボランティア",   18),
    ("農園の特徴・備考", 45),
]

# Write header row
for col_idx, (col_name, col_width) in enumerate(columns, start=1):
    cell = ws.cell(row=1, column=col_idx, value=col_name)
    cell.fill = header_fill
    cell.font = Font(bold=True, color="FFFFFF", size=11)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin
    ws.column_dimensions[get_column_letter(col_idx)].width = col_width

ws.row_dimensions[1].height = 30

# --- Data ---
data = [
    (
        "宮野園",
        "埼玉県",
        "狭山市大字北入曽25-2",
        "042-959-3025",
        "",
        "https://miyanoen.com/",
        "約1時間（狭山市）",
        "◎ 自社焙煎あり\n「狭山火入れ」独自焙煎\nほうじ茶販売あり",
        "茶摘み体験\n手作り茶体験\n抹茶アート\n茶の葉天ぷら試食",
        "要問い合わせ",
        "1852年創業。日本三大銘茶「狭山茶」の老舗。\n埼玉県認定「彩の国優良ブランド品」。\n体験プログラム充実。",
    ),
    (
        "狭山茶農家 ささら屋",
        "埼玉県",
        "入間市宮寺（詳細はサイト参照）",
        "（サイト参照）",
        "（サイト参照）",
        "https://www.sasaraya.com/",
        "約1時間（入間市）",
        "○ 茎ほうじ茶あり\n自園製茶のため焙煎対応\n可能性あり（要確認）",
        "茶摘み体験\n各種お茶講座",
        "要問い合わせ\n(SNSからも可)",
        "江戸時代から続く茶農家を2018年に刷新。\n「狭山茶はじまりの地」宮寺で活動。\nInstagram: @sasaraya.tea.farmer",
    ),
    (
        "和田園",
        "埼玉県",
        "所沢市三ケ島3-1157-3",
        "04-2948-0626",
        "",
        "https://wadaen.jp/",
        "約1時間（所沢市）",
        "△ 要確認\n自社仕上げ加工あり",
        "茶摘み体験（無料）\n工場見学\n新茶試飲\n※要電話予約",
        "要問い合わせ",
        "農林水産大臣賞を3度受賞。\n昭和10年茶栽培開始、1951年製茶所創業。\n減農薬・有機質肥料使用、エコファーマー認定。",
    ),
    (
        "浅見園",
        "埼玉県",
        "狭山市笹井2695-3",
        "04-2953-3790",
        "",
        "https://asamien.jp/",
        "約1時間（狭山市）",
        "◎ 自社焙煎あり\n「狭山火入れ」焙煎技術\nほうじ茶対応",
        "茶摘み体験\n製造・工場見学\n茶道体験（煎茶・抹茶）\n茶室「慈風庵」使用",
        "要問い合わせ",
        "1930年創業。製造から販売まで一貫体制。\n狭山火入れで甘くて濃厚な風味を引き出す。\n体験プログラム総合的に充実。",
    ),
    (
        "増岡園",
        "埼玉県",
        "入間市上谷ケ貫551-1",
        "04-2936-0250",
        "",
        "https://masuokaen.jp/",
        "約1時間（入間市）",
        "△ 要確認\n自社工場で製茶・仕上げあり",
        "工場見学\n茶摘み体験\n手もみ茶体験\n※要予約",
        "要問い合わせ",
        "狭山茶で唯一の有機認証工場。\n農薬・化学肥料不使用で30年以上の有機栽培。\n自家製和紅茶も製造。FAX: 04-2936-5858",
    ),
    (
        "神奈川県農協茶業センター\n（足柄茶）",
        "神奈川県",
        "足柄上郡山北町川西691-7",
        "0465-77-2001",
        "info@ashigaracha.co.jp",
        "https://www.ashigaracha.co.jp/",
        "約1.5時間（山北町）",
        "◎ ほうじ茶製品あり\nほうじ茶体験ワーク\nショップも実施\nフライパン焙煎体験可",
        "茶摘み体験\nほうじ茶作り体験\n新芽天ぷら試食\n茶娘衣装着付け\nファームオーナー制度\n（年会費16,500円）",
        "要問い合わせ\nFAX: 0465-77-2006",
        "足柄茶の中核施設。丹沢・足柄山麓の広域農家を束ねるJA茶業センター。\n南足柄市・小田原市・山北町など複数市町村が産地。",
    ),
]

for row_idx, row_data in enumerate(data, start=2):
    fill = row_odd_fill if row_idx % 2 == 1 else row_even_fill
    for col_idx, value in enumerate(row_data, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.border = thin
        cell.font = Font(size=10)
    ws.row_dimensions[row_idx].height = 75

# Freeze header row
ws.freeze_panes = "A2"

# Add a note row at the bottom
note_row = len(data) + 3
ws.cell(row=note_row, column=1,
        value="※ 調査日: 2026年4月19日  ／  ボランティア・援農の受け入れ可否は各農園に直接お問い合わせください。  ／  茶摘みシーズン（4〜5月）に問い合わせると情報が得やすいです。")
ws.cell(row=note_row, column=1).font = Font(italic=True, color="555555", size=9)
ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=len(columns))

output_path = "/home/user/Claude-Code-Test/東京近郊お茶農園リスト.xlsx"
wb.save(output_path)
print(f"Saved: {output_path}")
